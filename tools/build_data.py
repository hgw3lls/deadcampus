#!/usr/bin/env python3
"""Build public atlas data from the local archive sources."""

from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit(
        "openpyxl is required to rebuild atlas data. Use the Codex bundled Python "
        "runtime or install openpyxl in your local Python environment."
    ) from exc


ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = Path(__file__).resolve().parents[1] / "dist" / "data"
OUT_JS = DATA_DIR / "atlas-data.js"
OUT_JSON = DATA_DIR / "canonical-atlas.json"


TYPOLOGY = {
    "Absorbed": {
        "tone": "#355c7d",
        "description": "A closed campus taken into a larger educational system.",
    },
    "Parcelized": {
        "tone": "#a44732",
        "description": "A campus divided across buyers, agencies, or redevelopment parcels.",
    },
    "Speculative": {
        "tone": "#8a6f2d",
        "description": "Land held, marketed, auctioned, or repositioned as real estate.",
    },
    "Security": {
        "tone": "#4b4f37",
        "description": "Educational land converted to state, training, or security use.",
    },
    "Health": {
        "tone": "#2f6654",
        "description": "Institutional land reused by medical, hospital, or care systems.",
    },
    "Housing": {
        "tone": "#7a3f61",
        "description": "Campus property converted to residential or senior living use.",
    },
    "Educational reuse": {
        "tone": "#4d6b3c",
        "description": "The site remains educational under a new owner, mission, or scale.",
    },
    "K12": {
        "tone": "#b15f28",
        "description": "Higher education property converted to primary or secondary schooling.",
    },
    "Religious/K12": {
        "tone": "#6f4a7b",
        "description": "A mixed religious and school reuse of former college land.",
    },
    "Arts reuse": {
        "tone": "#286b74",
        "description": "Campus property reused by arts, music, or cultural programs.",
    },
    "Corporate": {
        "tone": "#574c41",
        "description": "Campus land or infrastructure absorbed into corporate use.",
    },
    "Cloud": {
        "tone": "#1f6875",
        "description": "Physical campus infrastructure converted into data infrastructure.",
    },
    "Ruin": {
        "tone": "#5c5650",
        "description": "The site remains abandoned, decayed, or unresolved.",
    },
}


PUBLIC_SUMMARIES = {
    ("birmingham southern college", "AL"): "After closing in 2024, the Birmingham campus moved from private liberal arts college to federal training infrastructure, with the U.S. Coast Guard opening TRACEN Birmingham-Southern in 2026.",
    ("green mountain college", "VT"): "Closed in 2019, the Poultney campus was bought at auction by private owner Raj Bhakta in 2020; by 2026, his redevelopment plan had stalled and the site was again being offered or repositioned.",
    ("notre dame college", "OH"): "Notre Dame College closed in 2024. In 2026, Akron Children's Hospital emerged as the auction winner for the South Euclid campus, pointing the site toward pediatric specialty care.",
    ("college of saint rose", "NY"): "The College of Saint Rose closed in 2024. Albany County's Pine Hills Land Authority then moved to acquire and redevelop the Albany campus as a public authority project.",
    ("wells college", "NY"): "Wells closed in 2024. A proposed Hiawatha Institute Indigenous college reuse became disputed in 2026 after Wells said the sale would not move forward, leaving the campus status unresolved.",
    ("fontbonne university", "MO"): "Fontbonne wound down instruction through summer 2025. Washington University acquired the 16-acre Clayton campus and leased it back during the closure runway, with long-term use still unsettled.",
    ("cabrini university", "PA"): "Cabrini closed after the 2023-24 academic year. Villanova assumed ownership of the 112-acre Radnor property and began converting it into Villanova University Cabrini Campus.",
    ("medaille university", "NY"): "Medaille closed in 2023. The Buffalo campus sold to Charter Schools Development Corporation for Buffalo Academy of Science Charter School, turning university buildings into a charter-school campus.",
    ("cardinal stritch university", "WI"): "The Fox Point campus is classified here as a religious and K-12 reuse after Cardinal Stritch's 2023 closure, moving Catholic higher-education property into private school conversion.",
    ("cazenovia college", "NY"): "After Cazenovia closed in 2023, its campus fragmented into mixed redevelopment, municipal ownership, and private reuse instead of remaining one college property.",
    ("presentation college", "SD"): "Presentation College's Aberdeen campus is tracked as a parcelized afterlife, with municipal and other buyers dividing a former Catholic health-oriented college footprint.",
    ("finlandia university", "MI"): "Finlandia's 2023 closure produced a civic and nonprofit fragmentation case, with pieces of the Hancock campus absorbed by various local groups.",
    ("mount ida college", "MA"): "Mount Ida's 2018 closure converted the Newton campus into a public university asset when UMass Amherst took control.",
    ("pine manor college", "MA"): "Pine Manor's 2020 merger folded the Brookline campus into Boston College, marking an elite institutional absorption.",
    ("wheelock college", "MA"): "Wheelock's 2018 merger with Boston University turned a small specialized college into part of a larger urban university system.",
    ("newbury college", "MA"): "After Newbury closed in 2019, Welltower moved the Brookline campus toward senior living, shifting educational property into housing infrastructure.",
    ("marlboro college", "VT"): "Marlboro's 2020 closure left the Vermont campus to an arts and music reuse under Marlboro Music.",
    ("burlington college", "VT"): "Burlington College's 2016 collapse led to property liquidation and developer or bank control, making the campus a speculative real estate case.",
    ("marygrove college", "MI"): "Marygrove's 2019 closure produced a community educational redevelopment model led by the Marygrove Conservancy in Detroit.",
    ("dowling college", "NY"): "Dowling's 2016 closure sent the Oakdale estate through auction and redevelopment, parcelizing a historic campus.",
    ("urbana university", "OH"): "Urbana closed in 2020 and entered private redevelopment marketing, a speculative afterlife for a regional campus.",
    ("atlantic union college", "MA"): "Atlantic Union's South Lancaster campus is tracked as slow parcel liquidation under religious ownership after its 2018 closure marker.",
    ("concordia university portland", "OR"): "After Concordia University Portland closed in 2020, the University of Oregon acquired the campus for public university use.",
    ("holy names university", "CA"): "Holy Names closed in 2023 and sold its Oakland hills campus to BH Properties, moving the site into investment ownership and leasing or redevelopment scenarios.",
    ("marymount california university", "CA"): "After Marymount California ceased operations in 2022, UCLA acquired its Rancho Palos Verdes and San Pedro sites to expand teaching, learning, and housing capacity.",
    ("macmurray college", "IL"): "MacMurray's 2020 closure culminated in a campus auction split into 18 tracts, a parcelized end to a 174-year college.",
    ("becker college", "MA"): "Becker's 2021 closure produced a split municipal and housing afterlife: Leicester acquired campus properties, with some buildings later used for emergency family housing.",
    ("college of new rochelle", "NY"): "The College of New Rochelle entered bankruptcy and auction after its 2019 closure, leaving an urban campus fragmented across buyers.",
    ("alderson broaddus university", "WV"): "Alderson Broaddus closed in 2023 and moved through bankruptcy sale to CGP Construction and Battlers Knob, shifting the Philippi campus into private redevelopment and community use.",
    ("daniel webster college", "NH"): "After the ITT-linked 2017 closure, Daniel Webster's Nashua campus went through auction sale to Xinhua Education, a speculative institutional reuse case.",
    ("virginia intermont college", "VA"): "Virginia Intermont's Bristol campus is tracked as a ruin and absentee ownership case after the 2014 closure, with the physical site unresolved.",
    ("southern vermont college", "VT"): "Southern Vermont College's 2019 closure led to health and resort redevelopment under Southwestern Vermont Health Care, repurposing the Bennington campus beyond higher education.",
    ("st gregory s university", "OK"): "St. Gregory's 2017 closure moved the Shawnee campus into corporate ownership under Hobby Lobby, marking a private institutional acquisition.",
    ("gwu virginia science and technology campus", "VA"): "GW's Ashburn campus was sold in 2026 to Amazon Data Services; academic programs can remain temporarily while the land enters cloud-infrastructure redevelopment pressure.",
}


PUBLIC_FIELD_OVERRIDES = {
    ("gwu virginia science and technology campus", "VA"): {
        "verification": "Non-closure asset sale",
    },
    ("wells college", "NY"): {
        "afterlife": "Disputed/pending Indigenous educational reuse",
        "reuseType": "Proposed Indigenous educational reuse",
        "buyerController": "Hiawatha Institute (sale disputed)",
        "verification": "Sale disputed",
        "transformation": "Proposed Indigenous educational reuse under contested sale",
    },
}


def rows_from_sheet(path: Path, sheet_name: str) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    output = []
    for row in rows[1:]:
        item = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            item[header] = row[index] if index < len(row) else None
        if any(value not in (None, "") for value in item.values()):
            output.append(item)
    return output


def csv_rows(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8-sig") as file:
        return list(csv.DictReader(file))


def norm(value: object) -> str:
    text = str(value or "").lower()
    text = text.replace("&", " and ")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def slug(value: str) -> str:
    text = norm(value)
    return re.sub(r"\s+", "-", text).strip("-")


def clean_year(value: object) -> int | None:
    if value in (None, ""):
        return None
    match = re.search(r"\d{4}", str(value))
    return int(match.group(0)) if match else None


def clean_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def canonical_morphology(value: object) -> str:
    raw = clean_text(value)
    if not raw:
        return ""
    fixed = raw.replace("_", " ").title()
    aliases = {
        "Arts Reuse": "Arts reuse",
        "Educational Reuse": "Educational reuse",
        "Branch / Satellite Retreat": "BRANCH / SATELLITE RETREAT",
        "Main / Root Closure": "MAIN / ROOT CLOSURE",
        "Distributed Site / Annex": "DISTRIBUTED SITE / ANNEX",
        "For-Profit System Collapse": "FOR-PROFIT SYSTEM COLLAPSE",
        "K12": "K12",
    }
    return aliases.get(fixed, fixed)


def is_url(value: object) -> bool:
    text = clean_text(value)
    return text.startswith("http://") or text.startswith("https://")


def choose(*values: object) -> str:
    for value in values:
        text = clean_text(value)
        if text:
            return text
    return ""


def build_public_sites() -> list[dict]:
    master = rows_from_sheet(ROOT / "_DATA" / "dead_campus_master_table_v1.xlsx", "DeadCampusAtlas")
    phase2 = rows_from_sheet(ROOT / "_DATA" / "dead_campus_atlas_phase2.xlsx", "PHASE2_CLASSIFIED")
    spatial = rows_from_sheet(
        ROOT / "_DATA" / "dead_campus_atlas_phase3_spatialization.xlsx",
        "SPATIALIZED_CAMPUSES",
    )
    forensics = rows_from_sheet(
        ROOT / "_DATA" / "dead_campus_atlas_phase4_afterlife_forensics.xlsx",
        "AFTERLIFE_FORENSICS",
    )
    queue = csv_rows(ROOT / "dead-campus-queue-filtered.csv")

    phase2_by_key = {(norm(r.get("Institution")), clean_text(r.get("State")).upper()): r for r in phase2}
    spatial_by_key = {(norm(r.get("Institution")), clean_text(r.get("State")).upper()): r for r in spatial}
    forensics_by_key = {(norm(r.get("Institution")), clean_text(r.get("State")).upper()): r for r in forensics}

    queue_by_key: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for line, row in enumerate(queue, start=2):
        row["_line"] = line
        queue_by_key[(norm(row.get("name")), clean_text(row.get("state")).upper())].append(row)

    sites = []
    for index, row in enumerate(master, start=1):
        state = clean_text(row.get("State")).upper()
        key = (norm(row.get("Institution")), state)
        p2 = phase2_by_key.get(key, {})
        sp = spatial_by_key.get(key, {})
        af = forensics_by_key.get(key, {})
        queue_rows = queue_by_key.get(key, [])
        year = clean_year(choose(row.get("Closure_Year"), p2.get("Closure_Year"), sp.get("Closure_Year")))
        name = choose(row.get("Institution"), p2.get("Institution"), sp.get("Institution"))
        morphology = canonical_morphology(choose(row.get("Morphology"), p2.get("Morphology"), sp.get("Morphology")))
        public_summary = PUBLIC_SUMMARIES.get(key, "")
        evidence = []
        for q in queue_rows:
            evidence.append(
                {
                    "line": q.get("_line"),
                    "sourceSheet": clean_text(q.get("sourceSheet")),
                    "verification": clean_text(q.get("verification")),
                    "sourceUrl": clean_text(q.get("sourceUrl")) if is_url(q.get("sourceUrl")) else "",
                    "sourceNote": "" if is_url(q.get("sourceUrl")) else clean_text(q.get("sourceUrl")),
                    "missingFields": clean_text(q.get("missingFields")),
                }
            )

        sites.append(
            {
                "id": f"{slug(name)}-{state.lower()}-{year or 'undated'}",
                "recordNumber": index,
                "name": name,
                "city": choose(sp.get("City")),
                "state": state,
                "region": choose(sp.get("Region"), p2.get("Region")),
                "closureYear": year,
                "sector": choose(row.get("Sector"), sp.get("Sector"), p2.get("Sector")),
                "campusType": choose(row.get("Campus_Type"), sp.get("Campus_Type"), p2.get("Campus_Type")),
                "status": choose(row.get("Status")),
                "reuseType": choose(row.get("Reuse_Type"), af.get("Afterlife")),
                "buyerController": choose(
                    row.get("Buyer_or_Controller"),
                    sp.get("Buyer_or_Controller"),
                    p2.get("Buyer_or_Controller"),
                    af.get("Buyer"),
                ),
                "buyerType": choose(af.get("Buyer_Type")),
                "morphology": morphology,
                "scale": choose(row.get("Scale")),
                "notes": choose(row.get("Notes")),
                "verification": choose(row.get("Verification"), "Confirmed"),
                "transformation": choose(sp.get("Transformation"), p2.get("Transformation"), af.get("Transformation_Logic")),
                "atlasPriority": choose(sp.get("Atlas_Priority"), p2.get("Atlas_Priority"), af.get("Atlas_Importance")),
                "salePrice": choose(af.get("Sale_Price")),
                "afterlife": choose(af.get("Afterlife"), row.get("Reuse_Type")),
                "latitude": sp.get("Latitude") if sp.get("Latitude") not in ("", None) else None,
                "longitude": sp.get("Longitude") if sp.get("Longitude") not in ("", None) else None,
                "publicSummary": public_summary,
                "evidence": evidence,
                "sourceFiles": [
                    "_DATA/dead_campus_master_table_v1.xlsx",
                    "_DATA/dead_campus_atlas_phase2.xlsx",
                    "_DATA/dead_campus_atlas_phase3_spatialization.xlsx",
                    "_DATA/dead_campus_atlas_phase4_afterlife_forensics.xlsx",
                    "dead-campus-queue-filtered.csv",
                ],
            }
        )
        if key in PUBLIC_FIELD_OVERRIDES:
            sites[-1].update(PUBLIC_FIELD_OVERRIDES[key])
    return sites


def build_abyss_summary() -> dict:
    rows = csv_rows(ROOT / "_DATA" / "dead_campus_phase9_peps_normalized.csv")
    class_counts = Counter(clean_text(r.get("ATLAS_PRELIMINARY_CLASS")) or "Unclassified" for r in rows)
    site_type_counts = Counter(clean_text(r.get("ATLAS_SITE_TYPE")) or "Unclassified" for r in rows)
    state_counts = Counter(clean_text(r.get("ATLAS_STATE")) or "Foreign/unknown" for r in rows)
    year_counts = Counter(str(clean_year(r.get("ATLAS_CLOSURE_YEAR")) or "Unknown") for r in rows)

    def ordered_counts(counter: Counter, numeric: bool = False) -> list[dict]:
        if numeric:
            keys = sorted(counter, key=lambda item: int(item) if item.isdigit() else 9999)
        else:
            keys = sorted(counter, key=lambda item: (-counter[item], item))
        return [{"label": key, "count": int(counter[key])} for key in keys]

    return {
        "totalRecords": len(rows),
        "sourceUrl": "https://www.ed.gov/sites/ed/files/offices/OSFAP/PEPS/docs/closedschoolsearchreq.xlsx",
        "yearRange": [
            min(clean_year(r.get("ATLAS_CLOSURE_YEAR")) for r in rows if clean_year(r.get("ATLAS_CLOSURE_YEAR"))),
            max(clean_year(r.get("ATLAS_CLOSURE_YEAR")) for r in rows if clean_year(r.get("ATLAS_CLOSURE_YEAR"))),
        ],
        "siteTypeCounts": ordered_counts(site_type_counts),
        "classCounts": ordered_counts(class_counts),
        "stateCounts": ordered_counts(state_counts)[:20],
        "timeline": ordered_counts(year_counts, numeric=True),
        "missing": {
            "state": sum(1 for r in rows if not clean_text(r.get("ATLAS_STATE"))),
            "morphology": sum(1 for r in rows if not clean_text(r.get("Morphology"))),
            "afterlifeStatus": sum(1 for r in rows if not clean_text(r.get("Afterlife_Status"))),
        },
    }


def build_queue_audit() -> dict:
    rows = csv_rows(ROOT / "dead-campus-queue-filtered.csv")
    blank_counts = {}
    for field in [
        "closureYear",
        "state",
        "city",
        "sourceUrl",
        "verification",
        "buyerController",
        "afterlifeStatus",
    ]:
        blank_counts[field] = sum(1 for r in rows if not clean_text(r.get(field)))

    non_url_source = sum(1 for r in rows if clean_text(r.get("sourceUrl")) and not is_url(r.get("sourceUrl")))
    non_peps = [r for r in rows if clean_text(r.get("sourceSheet")) != "dead_campus_atlas__PEPS_SITE_"]
    duplicate_groups: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in non_peps:
        duplicate_groups[(norm(row.get("name")), clean_text(row.get("state")).upper())].append(row)
    duplicate_group_count = sum(1 for values in duplicate_groups.values() if len(values) > 1)
    duplicate_row_count = sum(len(values) for values in duplicate_groups.values() if len(values) > 1)

    return {
        "totalQueueRows": len(rows),
        "blankCounts": blank_counts,
        "nonUrlSourceUrlRows": non_url_source,
        "nonPepsRows": len(non_peps),
        "nonPepsDuplicateGroups": duplicate_group_count,
        "nonPepsDuplicateRows": duplicate_row_count,
        "statusCounts": [
            {"label": label, "count": int(count)}
            for label, count in Counter(clean_text(r.get("statusCode")) for r in rows).most_common()
        ],
    }


def main() -> None:
    sites = build_public_sites()
    morphologies = sorted({site["morphology"] for site in sites if site["morphology"]})
    payload = {
        "generatedFrom": [
            "dead-campus-queue-filtered.csv",
            "_DATA/dead_campus_master_table_v1.xlsx",
            "_DATA/dead_campus_atlas_phase2.xlsx",
            "_DATA/dead_campus_atlas_phase3_spatialization.xlsx",
            "_DATA/dead_campus_atlas_phase4_afterlife_forensics.xlsx",
            "_DATA/dead_campus_phase9_peps_normalized.csv",
        ],
        "publicSites": sites,
        "typology": {key: TYPOLOGY.get(key, {"tone": "#4f4a45", "description": ""}) for key in morphologies},
        "abyss": build_abyss_summary(),
        "queueAudit": build_queue_audit(),
    }
    OUT_JS.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    OUT_JS.write_text(
        "window.DEAD_CAMPUS_ATLAS = "
        + json.dumps(payload, ensure_ascii=False, indent=2)
        + ";\n",
        encoding="utf-8",
    )
    print(f"Wrote {OUT_JSON}")
    print(f"Wrote {OUT_JS}")


if __name__ == "__main__":
    main()
